import openpyxl, json, re, datetime

SRC = 'fin.xlsm'
wbv = openpyxl.load_workbook(SRC, data_only=True)
wbf = openpyxl.load_workbook(SRC)  # for fill colours

def num(v):
    if v is None or v == '':
        return 0
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    try:
        return round(float(str(v).replace(',', '')), 2)
    except Exception:
        return 0

def clean(s):
    return re.sub(r'\s+', ' ', str(s)).strip() if s is not None else ''

data = {
    'meta': {'version': 1, 'updated': datetime.date.today().isoformat(), 'source': 'כלכלת_משפחה.xlsm'},
    'expenses': {}, 'incomes': {}, 'budget': {},
}

# ---------- Yearly expense sheets (2024-2026) ----------
# columns D..O = Dec..Jan ; P = category ; Q = domain
MONTH_COLS = list(range(15, 3, -1))  # O(15)=Jan ... D(4)=Dec
for year in ['2024', '2025', '2026']:
    ws = wbv[year]; wf = wbf[year]
    groups, cur = [], None
    for r in range(3, ws.max_row + 1):
        name = clean(ws.cell(r, 16).value)
        dom = clean(ws.cell(r, 17).value)
        if name.startswith('סה"כ הוצאות'):
            break
        if dom:
            cur = {'domain': dom.replace('\n', ' '), 'items': []}
            groups.append(cur)
        if not name or name == 'סה"כ':
            continue
        if cur is None:
            cur = {'domain': 'כללי', 'items': []}; groups.append(cur)
        months = [num(ws.cell(r, c).value) for c in MONTH_COLS]
        col = [bool(wf.cell(r, c).fill and wf.cell(r, c).fill.fill_type) for c in MONTH_COLS]
        nz = [col[i] for i in range(12) if months[i]] or col
        fixed = sum(nz) * 2 > len(nz)  # Excel: coloured cell = fixed expense
        cur['items'].append({'name': name, 'fixed': fixed, 'months': months})
    groups = [g for g in groups if g['items']]
    data['expenses'][year] = groups
    # investments row
    for r in range(ws.max_row, 100, -1):
        if 'INVEST' in clean(ws.cell(r, 16).value):
            data['budget'][year] = {'investments': [num(ws.cell(r, c).value) for c in MONTH_COLS]}
            break

# ---------- Hidden monthly sheets (07-2022 .. 12-2023) ----------
for ws in wbv.worksheets:
    m = re.match(r'^(\d\d)-(\d{4})$', ws.title)
    if not m:
        continue
    wsf = wbf[ws.title]
    mi, year = int(m.group(1)) - 1, m.group(2)
    # which section-total rows are included in the grand total formula
    inc_rows = None
    gt = [r for r in range(1, ws.max_row + 1)
          if isinstance(wsf.cell(r, 7).value, str) and wsf.cell(r, 7).value.startswith('=SUM(') and ',' in wsf.cell(r, 7).value]
    for r in gt[-1:]:
        if True:
            f = str(wsf.cell(r, 7).value)
            inc_rows = set()
            for a, b in re.findall(r'G(\d+)(?::G(\d+))?', f):
                if b: inc_rows.update(range(int(a), int(b) + 1))
                else: inc_rows.add(int(a))
    groups = data['expenses'].setdefault(year, [])
    # first pass: section boundaries
    sections, cur = [], None
    for r in range(2, ws.max_row + 1):
        name = clean(ws.cell(r, 8).value); dom = clean(ws.cell(r, 9).value)
        if dom in ('סה"כ כללי', 'תחום') or (gt and r >= gt[-1]):
            continue
        if dom:
            cur = {'dom': dom, 'rows': []}; sections.append(cur)
        if cur is None: continue
        if name == 'סה"כ':
            cur['total'] = r; cur = None; continue
        if name: cur['rows'].append(r)
    for sec in sections:
        excluded = inc_rows is not None and sec.get('total') not in inc_rows
        dname = sec['dom'] + (' (לא נכלל בסיכום)' if excluded else '')
        g = next((g for g in groups if g['domain'] == dname), None)
        if g is None:
            g = {'domain': dname, 'items': []}
            if excluded: g['exclude'] = True
            groups.append(g)
        for r in sec['rows']:
            name = clean(ws.cell(r, 8).value)
            it = next((i for i in g['items'] if i['name'] == name), None)
            if it is None:
                it = {'name': name, 'fixed': False, 'months': [0] * 12}; g['items'].append(it)
            it['months'][mi] = num(ws.cell(r, 7).value)

# ---------- Incomes ----------
ws = wbv['הכנסות']
r = 1
while r <= ws.max_row:
    y = ws.cell(r, 4).value
    if isinstance(y, int) and 2000 < y < 2100 and ws.cell(r + 1, 3).value == 'מקור הכנסה':
        rows = []
        rr = r + 3
        while rr <= ws.max_row:
            n = clean(ws.cell(rr, 3).value)
            if n == 'סה"כ':
                break
            if n:
                rows.append({'name': n, 'months': [num(ws.cell(rr, c).value) for c in range(4, 16)]})
            rr += 1
        data['incomes'][str(y)] = rows
        r = rr
    r += 1

# ---------- Net worth (Total View) ----------
ws = wbv['Total View']
data['rates'] = {'usd': num(ws['P4'].value), 'eur': num(ws['O4'].value)}
assets = []
cur_map = {}
for r in range(28, 49):
    n = clean(ws.cell(r, 20).value)
    if not n:
        continue
    assets.append({
        'name': n,
        'value': num(ws.cell(r, 16).value),
        'pension': num(ws.cell(r, 15).value) or None,
        'buyPrice': num(ws.cell(r, 18).value) or None,
        'address': clean(ws.cell(r, 19).value) or None,
        'foreign': num(ws.cell(r, 21).value) or None,
        'updated': clean(ws.cell(r, 22).value) or None,
    })
# asset groups follow the Excel formulas
groups_def = [
    ('נדל"ן (דירות-בתים)', ['דירה-עיןהקורא34', 'דירה-בןגוריון33', 'דירה-ישראלבןציון7', 'פורטוגל - דירה1 (פורטו)', 'פורטוגל - דירה2 (גאיה)']),
    ('קרנות השתלמות', ['קרן השתלמות טל', 'קרן השתלמות רונן']),
    ('תיקי מסחר עצמאי', ['תיק השקעות IB חו"ל', 'תיק השקעות IBI ארץ']),
    ('קריפטו', ['קריפטו']),
    ('קרנות הון (חול)', ['השקעת קרן METAOR2', 'השקעת קרן METAOR6', 'השקעה קרן AUSTIN', 'השקעות קרן METAOR SPRINT']),
    ('פנסיות', ['פנסייה רונן', 'פנסייה טל']),
    ('עוש אמריקה', ['עוש בנק אמריקה']),
    ('קרנות כספית', ['כספית לאומי', 'כספית מזרחי']),
    ('מטח', ['סקונדרי מטח (לאומי)']),
]
for a in assets:
    a['group'] = next((g for g, names in groups_def if a['name'] in names), 'אחר')
# foreign-currency info (formula driven in Excel)
fx = {'פורטוגל - דירה1 (פורטו)': ('x4', 'פי 4 ממחיר הקנייה'), 'פורטוגל - דירה2 (גאיה)': ('x4', '')}
for a in assets:
    if a['name'] in ('קריפטו', 'תיק השקעות IB חו"ל', 'השקעת קרן METAOR2', 'השקעת קרן METAOR6',
                     'השקעה קרן AUSTIN', 'עוש בנק אמריקה', 'השקעות קרן METAOR SPRINT', 'סקונדרי מטח (לאומי)'):
        a['currency'] = 'USD'
    elif a['name'].startswith('פורטוגל'):
        a['currency'] = 'EUR_X4'   # value = foreign * 4  (as in Excel – to confirm)
    else:
        a['currency'] = 'ILS'
data['assets'] = assets
data['assetGroups'] = [g for g, _ in groups_def]
liab = []
for r in range(17, 25):
    n = clean(ws.cell(r, 17).value)
    if n:
        liab.append({'name': n, 'value': num(ws.cell(r, 16).value)})
# Porto mortgage is 127790 EUR in Excel formula
for l in liab:
    if l['name'] == 'משכנתא פורטו':
        l['foreign'] = 127790; l['currency'] = 'EUR'
data['liabilities'] = liab

# ---------- Net worth history ----------
ws = wbv['Total View - Graph']
hist = []
for r in range(2, ws.max_row + 1):
    d = ws.cell(r, 4).value
    if isinstance(d, datetime.datetime):
        hist.append({'date': d.date().isoformat(), 'assets': num(ws.cell(r, 1).value),
                     'liabilities': num(ws.cell(r, 2).value), 'net': num(ws.cell(r, 3).value)})
data['netWorthHistory'] = hist

# ---------- Graph sheet (historic mortgage / loans series) ----------
ws = wbv['גרפים']
gs = []
for r in range(3, ws.max_row + 1):
    d = ws.cell(r, 1).value
    if isinstance(d, datetime.datetime) and any(ws.cell(r, c).value for c in range(2, 8)):
        gs.append({'month': d.strftime('%Y-%m'), 'mashBenGurion': num(ws.cell(r, 4).value),
                   'mashEinHakore': num(ws.cell(r, 5).value), 'mashBatYam': num(ws.cell(r, 6).value),
                   'loans': num(ws.cell(r, 7).value)})
data['graphHistory'] = gs

# ---------- Risk management ----------
ws = wbv['ניהול סיכונים והגנות']
people = [('יובל', 8, 7), ('אופיר', 10, 9), ('מתן', 12, 11), ('טל', 14, 13), ('רונן', 16, 15)]
rows = []
for r in range(9, 19):
    lbl = clean(ws.cell(r, 17).value)
    if not lbl:
        continue
    row = {'label': lbl, 'people': {}}
    for p, vc, pc in people:
        v = ws.cell(r, vc).value; pr = ws.cell(r, pc).value
        if v is not None or pr is not None:
            row['people'][p] = {'value': v, 'premium': pr}
    rows.append(row)
data['risk'] = {
    'coverage': rows,
    'survivors': {
        'split': {'orphan': 0.4, 'tal': 0.6, 'ronen': 0.6},
        'ifRonenDies': {'orphan': num(ws['N22'].value), 'tal': num(ws['O22'].value)},
        'ifTalDies': {'orphan': num(ws['N23'].value), 'ronen': num(ws['P23'].value)},
    }
}

# ---------- Contacts ----------
ws = wbv['רשימת סוכנים ואנשי קשר']
data['contacts'] = [{'name': clean(ws.cell(r, 11).value), 'topic': clean(ws.cell(r, 10).value),
                     'phone': clean(ws.cell(r, 9).value), 'notes': clean(ws.cell(r, 8).value)}
                    for r in range(4, ws.max_row + 1) if ws.cell(r, 11).value]

# ---------- Inheritance projections ----------
ws = wbv['הורשה לילדים']
data['inheritance'] = {
    'growth': num(ws['H2'].value),
    'baseYear': 2026,
    'items': [
        {'key': 'ronenPen', 'name': 'פנסיה רונן', 'base': num(ws['J9'].value)},
        {'key': 'ronenTF', 'name': 'קרן השתלמות רונן', 'base': num(ws['K9'].value)},
        {'key': 'talPen', 'name': 'פנסיה טל', 'base': num(ws['M9'].value)},
        {'key': 'talTF', 'name': 'קרן השתלמות טל', 'base': num(ws['N9'].value)},
        {'key': 'ibi', 'name': 'תיק IBI ישראל', 'base': num(ws['P9'].value)},
        {'key': 'ibkr', 'name': 'תיק IBKR ארה"ב', 'base': num(ws['Q9'].value)},
    ],
    'history': [{'year': ws.cell(r, 7).value, 'values': [num(ws.cell(r, c).value) for c in (10, 11, 13, 14, 16, 17)]}
                for r in range(6, 9)],
    'endYear': 2042,
}

# ---------- Portugal ----------
ws = wbv['פורטוגל-נתונים']
pt = []
for r in range(6, ws.max_row + 1):
    if ws.cell(r, 8).value:
        amt = ws.cell(r, 9).value
        pt.append({'unit': clean(ws.cell(r, 8).value),
                   'amount': num(str(amt).replace('eu', '')) if amt is not None else 0,
                   'purpose': clean(ws.cell(r, 10).value), 'date': clean(ws.cell(r, 11).value)})
data['portugal'] = {'currency': 'EUR', 'transfers': pt}

data['settings'] = {
    'familyName': 'משפחת מוזס',
    'budgetSplit': {'need': 50, 'want': 30, 'invest': 20},
    'people': {'p1': 'רונן', 'p2': 'טל'},
}

json.dump(data, open('data.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

# ---- validation against Excel totals ----
for year, row in [('2026', 119), ('2025', 119), ('2024', 127)]:
    ws = wbv[year]
    tot = [0] * 12
    for g in data['expenses'][year]:
        for it in g['items']:
            for i, v in enumerate(it['months']):
                tot[i] += v
    xl = [num(ws.cell(row, c).value) for c in MONTH_COLS]
    fx = [0] * 12
    for g in data['expenses'][year]:
        for it in g['items']:
            if it['fixed']:
                for i, v in enumerate(it['months']): fx[i] += v
    xlf = [num(ws.cell(row - 2, c).value) for c in MONTH_COLS]
    print(year, 'total diff', [round(a - b) for a, b in zip(tot, xl)])
    print(year, 'fixed diff', [round(a - b) for a, b in zip(fx, xlf)])
for y in ['2022', '2023']:
    print(y, [round(sum(it['months'][i] for g in data['expenses'][y] if not g.get('exclude') for it in g['items'])) for i in range(12)])
